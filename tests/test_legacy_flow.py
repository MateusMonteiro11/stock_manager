import importlib
import os
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from contracts.inventory_snapshot import Marketplace
from services.auditoria_service import realizar_auditoria
from services.relatorio_service import gerar_relatorio_xlsx


class FakeProductRepository:
    def __init__(self):
        self.updated = None

    def find_by_id(self, product_id):
        return (product_id, "Torneira", "TOR-001", "Metais", 20)

    def update_stock(self, product_id, quantity):
        self.updated = (product_id, quantity)


class FakeAuditRepository:
    def __init__(self):
        self.created = None
        self.marketplace_stocks = []

    def create(self, consolidated_stock, status, audited_at, product_id):
        self.created = (consolidated_stock, status, audited_at, product_id)
        return 10

    def save_marketplace_stock(self, found_stock, audit_id, marketplace):
        self.marketplace_stocks.append((found_stock, audit_id, marketplace))


class FakeReportRepository:
    def list_all(self):
        return [(10, "Torneira", 16, "DIVERGENTE", "13/09/2026 22:00:00")]

    def list_marketplaces(self, audit_id):
        return [("Shopee", 16), ("Mercado Livre", 17)]


class LegacyFlowTests(unittest.TestCase):
    def test_auditoria_funciona_com_shopee_e_mercado_livre(self):
        products = FakeProductRepository()
        audits = FakeAuditRepository()

        success = realizar_auditoria(
            1,
            16,
            17,
            product_repository=products,
            audit_repository=audits,
        )

        self.assertTrue(success)
        self.assertEqual(audits.created[0:2], (16, "DIVERGENTE"))
        self.assertEqual(products.updated, (1, 16))
        self.assertEqual(
            {entry[2] for entry in audits.marketplace_stocks},
            {Marketplace.SHOPEE, Marketplace.MERCADO_LIVRE},
        )
        self.assertEqual(len(audits.marketplace_stocks), 2)

    def test_relatorio_xlsx_preserva_os_dois_marketplaces(self):
        with TemporaryDirectory() as directory:
            report_path = Path(directory) / "auditoria.xlsx"

            result = gerar_relatorio_xlsx(
                report_path,
                repository=FakeReportRepository(),
            )

            workbook = load_workbook(result)
            rows = list(workbook["Auditorias"].iter_rows(values_only=True))
            self.assertEqual(rows[1][5], "Shopee")
            self.assertEqual(rows[2][5], "Mercado Livre")
            self.assertEqual(len(rows), 3)

    def test_repository_sqlite_resolve_marketplace_sem_id_fixo(self):
        with TemporaryDirectory() as directory:
            previous_path = os.environ.get("MYSTOCK_SQLITE_DB_PATH")
            os.environ["MYSTOCK_SQLITE_DB_PATH"] = str(Path(directory) / "estoque.db")
            module_names = [
                "connection.auditoria_repository",
                "connection.conexao",
            ]
            for module_name in module_names:
                sys.modules.pop(module_name, None)

            try:
                conexao = importlib.import_module("connection.conexao")
                repository = importlib.import_module(
                    "connection.auditoria_repository"
                )
                from connection.init_db import inicializar_banco

                inicializar_banco(conexao.conn)
                conexao.cursor.execute(
                    """
                    INSERT INTO produtos (nome, sku, categoria)
                    VALUES (?, ?, ?)
                    """,
                    ("Torneira", "TOR-001", "Metais"),
                )
                product_id = conexao.cursor.lastrowid
                audit_id = repository.criar_auditoria(
                    16,
                    "DIVERGENTE",
                    "13/09/2026 22:00:00",
                    product_id,
                )

                self.assertTrue(
                    repository.salvar_estoque_marketplace(
                        16,
                        audit_id,
                        Marketplace.SHOPEE,
                    )
                )
                self.assertTrue(
                    repository.salvar_estoque_marketplace(
                        17,
                        audit_id,
                        Marketplace.MERCADO_LIVRE,
                    )
                )
                self.assertEqual(
                    set(repository.listar_marketplaces_auditoria(audit_id)),
                    {("Shopee", 16), ("Mercado Livre", 17)},
                )
            finally:
                if "conexao" in locals():
                    conexao.conn.close()
                for module_name in module_names:
                    sys.modules.pop(module_name, None)
                if previous_path is None:
                    os.environ.pop("MYSTOCK_SQLITE_DB_PATH", None)
                else:
                    os.environ["MYSTOCK_SQLITE_DB_PATH"] = previous_path


if __name__ == "__main__":
    unittest.main()
